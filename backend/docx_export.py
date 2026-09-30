"""Basic DOCX export for documents."""
import io
from docx import Document as Docx
from docx.shared import Pt, RGBColor
from renderer import compute_totals, fmt_money, TYPE_META


def _set_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = RGBColor(*color)


ACCENT = (0xC0, 0x4C, 0x20)
MUTED = (0x78, 0x74, 0x6C)


def build_docx(doc: dict) -> bytes:
    d = doc.get("data", {}) or {}
    currency = doc.get("currency", "USD")
    dtype = doc.get("type")
    meta = TYPE_META.get(dtype, {})
    layout = meta.get("layout", "content")
    docx = Docx()

    # Header
    h = docx.add_paragraph()
    r = h.add_run(d.get("brand", "Your Agency"))
    _set_font(r, 16, True)
    if d.get("tagline"):
        t = docx.add_paragraph()
        _set_font(t.add_run(d["tagline"]), 9, color=MUTED)

    title = docx.add_paragraph()
    _set_font(title.add_run(f"{d.get('label', meta.get('label',''))}  "), 24, True)
    _set_font(title.add_run(d.get("number", "")), 12, color=ACCENT)

    if layout == "financial":
        p = docx.add_paragraph()
        _set_font(p.add_run(f"Bill to: {d.get('bill_to_name','')}"), 11, True)
        docx.add_paragraph(d.get("bill_to_lines", ""))
        docx.add_paragraph(f"From: {d.get('from_name','')}")
        docx.add_paragraph(f"Issue date: {d.get('issue_date','')}    Due date: {d.get('due_date','')}")
        if d.get("project_reference"):
            docx.add_paragraph(f"Reference: {d.get('project_reference')}")
        items = doc.get("line_items", []) or []
        table = docx.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, htxt in enumerate(["Description", "Qty", "Rate", "Amount"]):
            _set_font(hdr[i].paragraphs[0].add_run(htxt), 10, True)
        for it in items:
            amt = float(it.get("qty", 0) or 0) * float(it.get("rate", 0) or 0)
            cells = table.add_row().cells
            cells[0].text = str(it.get("description", ""))
            if it.get("sub"):
                _set_font(cells[0].add_paragraph().add_run(it["sub"]), 8, color=MUTED)
            cells[1].text = str(it.get("qty", ""))
            cells[2].text = fmt_money(it.get("rate", 0), currency)
            cells[3].text = fmt_money(amt, currency)
        t = compute_totals(doc)
        docx.add_paragraph(f"Subtotal: {fmt_money(t['subtotal'], currency)}")
        if (doc.get("discount") or {}).get("enabled"):
            docx.add_paragraph(f"Discount: -{fmt_money(t['discount'], currency)}")
        if (doc.get("tax") or {}).get("enabled"):
            docx.add_paragraph(f"Tax: {fmt_money(t['tax'], currency)}")
        tot = docx.add_paragraph()
        _set_font(tot.add_run(f"Total: {fmt_money(t['total'], currency)}"), 14, True, ACCENT)
        docx.add_paragraph(f"Payment details: {d.get('payment_details','')}")
        docx.add_paragraph(f"Payment terms: {d.get('payment_terms','')}")
    elif layout == "letterhead":
        docx.add_paragraph(f"Date: {d.get('issue_date','')}")
        docx.add_paragraph(d.get("letter_body", ""))
    else:
        docx.add_paragraph(f"Prepared for: {d.get('bill_to_name','')}    Date: {d.get('issue_date','')}")
        for s in (d.get("sections", []) or []):
            hp = docx.add_paragraph()
            _set_font(hp.add_run(s.get("heading", "")), 13, True)
            docx.add_paragraph(s.get("body", ""))
        if (doc.get("line_items") or []):
            t = compute_totals(doc)
            tot = docx.add_paragraph()
            _set_font(tot.add_run(f"Total: {fmt_money(t['total'], currency)}"), 13, True, ACCENT)
        if dtype in ("service_agreement", "nda", "statement_of_work"):
            docx.add_paragraph("Note: Legal text must be reviewed by a qualified lawyer before use.")

    docx.add_paragraph(d.get("footer_line", "Thank you for working with us."))
    buf = io.BytesIO()
    docx.save(buf)
    return buf.getvalue()
